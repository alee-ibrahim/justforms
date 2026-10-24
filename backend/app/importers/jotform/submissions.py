"""Jotform submission exports (.xlsx / .csv) -> JustForms answers (DESIGN 13.3).

``read_table`` turns the file into a header row and data rows; ``suggest_mapping`` matches
columns to questions by label; ``convert_row`` turns one row into answers, metadata, links for
file and signature columns, and per-cell problems.
"""

from __future__ import annotations

import csv
import io
import re
import unicodedata
from collections.abc import Iterator
from dataclasses import dataclass, field
from datetime import date, datetime, time, tzinfo
from typing import Any

from app.logic.index import FormIndex
from app.schemas.form import (
    DateField,
    EmailField,
    FieldElement,
    FileField,
    Form,
    HiddenField,
    LongTextField,
    MultiChoiceField,
    NumberField,
    PhoneField,
    RatingField,
    ShortTextField,
    SignatureField,
    SingleChoiceField,
    TimeField,
    YesNoField,
)
from app.validation.answers import check_answer

MAX_FILE_BYTES = 50 * 1024 * 1024
MAX_ROWS = 100_000
MAX_COLUMNS = 1000
SAMPLE_ROWS = 20

META_TARGETS = {
    "meta:submitted_at": "Submission date",
    "meta:external_id": "Jotform submission ID",
    "meta:ip": "IP address",
}
IGNORE = "ignore"
_META_HEADERS = {
    "submission date": "meta:submitted_at",
    "submitted at": "meta:submitted_at",
    "submission id": "meta:external_id",
    "submissionid": "meta:external_id",
    "ip": "meta:ip",
    "ip address": "meta:ip",
}
YES = {"yes", "true", "1", "y", "އާން"}
NO = {"no", "false", "0", "n", "ނޫން"}
# Problems that historic answers may have (relative date limits moved on since).
_ACCEPTED = {"date_from", "date_to"}


class TableError(ValueError):
    """``code``: ``empty`` | ``unreadable`` | ``too_many_rows`` | ``too_many_columns``."""

    def __init__(self, code: str) -> None:
        super().__init__(code)
        self.code = code


@dataclass
class Table:
    headers: list[str]
    rows: list[list[Any]]


def _cell(value: Any) -> Any:
    if isinstance(value, str):
        return value.strip()
    return value


def _rows_csv(data: bytes) -> Iterator[list[Any]]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("cp1252", "replace")
    sample = text[:20_000]
    try:
        dialect: Any = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    yield from csv.reader(io.StringIO(text, newline=""), dialect)


def _rows_xlsx(data: bytes) -> Iterator[list[Any]]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        for row in ws.iter_rows(values_only=True):
            yield list(row)
    finally:
        wb.close()


def read_table(data: bytes, filename: str) -> Table:
    is_xlsx = data[:2] == b"PK" or filename.lower().endswith((".xlsx", ".xlsm"))
    try:
        rows_iter = _rows_xlsx(data) if is_xlsx else _rows_csv(data)
        headers: list[str] | None = None
        rows: list[list[Any]] = []
        for raw in rows_iter:
            values = [_cell(v) for v in raw]
            if headers is None:
                if not any(v not in (None, "") for v in values):
                    continue
                while values and values[-1] in (None, ""):
                    values.pop()
                if len(values) > MAX_COLUMNS:
                    raise TableError("too_many_columns")
                headers = ["" if v is None else str(v) for v in values]
                continue
            if not any(v not in (None, "") for v in values):
                continue
            rows.append((values + [None] * len(headers))[: len(headers)])
            if len(rows) > MAX_ROWS:
                raise TableError("too_many_rows")
    except TableError:
        raise
    except Exception as exc:  # malformed workbook/CSV: report, never crash
        raise TableError("unreadable") from exc
    if not headers:
        raise TableError("empty")
    return Table(headers=headers, rows=rows)


# -- mapping -------------------------------------------------------------------------------


def norm(text: str) -> str:
    text = unicodedata.normalize("NFKC", str(text)).replace("\u00a0", " ").lower()
    return re.sub(r"[\W_]+", " ", text).strip()


def importable_fields(form: Form) -> list[FieldElement]:
    """Top-level questions in form order (repeater rows cannot come from a flat export)."""
    out: list[Any] = []
    for info in FormIndex.build(form).elements.values():
        if info.kind == "field" and info.repeater_id is None:
            out.append(info.element)
    return out


def targets(form: Form) -> list[dict[str, str]]:
    out = [
        {"value": f"field:{f.id}", "label": f.label or f.id, "type": f.type}
        for f in importable_fields(form)
    ]
    out += [{"value": k, "label": v, "type": "meta"} for k, v in META_TARGETS.items()]
    return out


def suggest_mapping(headers: list[str], form: Form) -> list[str]:
    """One target per column. Repeated labels (three "Name" columns) map in order onto the
    questions with that label."""
    fields = importable_fields(form)
    by_label: dict[str, list[str]] = {}
    for f in fields:
        by_label.setdefault(norm(f.label), []).append(f.id)
    by_id = {norm(f.id): f.id for f in fields}
    used: set[str] = set()
    out: list[str] = []
    for header in headers:
        key = norm(header)
        meta = _META_HEADERS.get(key)
        if meta and meta not in used:
            out.append(meta)
            used.add(meta)
            continue
        choice = next((fid for fid in by_label.get(key, []) if f"field:{fid}" not in used), None)
        if choice is None and key in by_id and f"field:{by_id[key]}" not in used:
            choice = by_id[key]
        if choice is None:
            out.append(IGNORE)
            continue
        used.add(f"field:{choice}")
        out.append(f"field:{choice}")
    return out


# -- conversion ----------------------------------------------------------------------------

_MONTHS = {
    m: i
    for i, names in enumerate(
        [
            ("jan", "january"),
            ("feb", "february"),
            ("mar", "march"),
            ("apr", "april"),
            ("may",),
            ("jun", "june"),
            ("jul", "july"),
            ("aug", "august"),
            ("sep", "sept", "september"),
            ("oct", "october"),
            ("nov", "november"),
            ("dec", "december"),
        ],
        start=1,
    )
    for m in names
}


def parse_date(value: Any, order: str = "auto") -> date | None:
    """Jotform date formats: ``2026-10-24``, ``2026/10/24``, ``10/24/2026`` or ``24/10/2026``
    (``order`` mdy/dmy; auto picks dmy when the first number is above 12), ``24-10-2026``,
    ``Oct 24, 2026``, ``24 Oct 2026``; times after the date are ignored."""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    if not text:
        return None
    try:
        m = re.match(r"^(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})(?!\d)", text)
        if m:
            return date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
        m = re.match(r"^(\d{1,2})[-/.](\d{1,2})[-/.](\d{4})(?!\d)", text)
        if m:
            a, b, year = int(m.group(1)), int(m.group(2)), int(m.group(3))
            dmy = order == "dmy" or (order == "auto" and a > 12)
            if text[len(m.group(1))] == "-" and order == "auto":
                dmy = a > 12 or b <= 12  # Jotform's DD-MM-YYYY uses dashes
            return date(year, b, a) if dmy else date(year, a, b)
        m = re.match(r"^([A-Za-z]{3,9})\.? (\d{1,2}),? (\d{4})\b", text)
        if m and m.group(1).lower() in _MONTHS:
            return date(int(m.group(3)), _MONTHS[m.group(1).lower()], int(m.group(2)))
        m = re.match(r"^(\d{1,2}) ([A-Za-z]{3,9})\.?,? (\d{4})\b", text)
        if m and m.group(2).lower() in _MONTHS:
            return date(int(m.group(3)), _MONTHS[m.group(2).lower()], int(m.group(1)))
    except ValueError:
        return None
    return None


def parse_time(value: Any) -> str | None:
    if isinstance(value, datetime):
        return value.strftime("%H:%M")
    if isinstance(value, time):
        return value.strftime("%H:%M")
    text = str(value or "").strip()
    m = re.search(r"(?<!\d)(\d{1,2}):(\d{2})(?::\d{2})?(?:\.\d+)?\s*([AaPp][Mm])?\b", text)
    if not m:
        return None
    hour, minute = int(m.group(1)), int(m.group(2))
    if m.group(3):
        hour = hour % 12 + (12 if m.group(3).lower() == "pm" else 0)
    if hour > 23 or minute > 59:
        return None
    return f"{hour:02d}:{minute:02d}"


def parse_datetime(value: Any, tz: tzinfo, order: str = "auto") -> datetime | None:
    if isinstance(value, datetime):
        return value if value.tzinfo else value.replace(tzinfo=tz)
    day = parse_date(value, order)
    if day is None:
        return None
    clock = parse_time(value) or "00:00"
    hour, minute = (int(x) for x in clock.split(":"))
    return datetime(day.year, day.month, day.day, hour, minute, tzinfo=tz)


def parse_number(value: Any) -> float | int | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, int | float):
        return value
    text = str(value or "").strip().replace(" ", "")
    if not text:
        return None
    if "," in text and "." in text:
        text = (
            text.replace(",", "")
            if text.rfind(".") > text.rfind(",")
            else (text.replace(".", "").replace(",", "."))
        )
    elif text.count(",") == 1 and len(text.split(",")[1]) != 3:
        text = text.replace(",", ".")
    else:
        text = text.replace(",", "")
    try:
        number = float(text)
    except ValueError:
        return None
    return int(number) if number.is_integer() else number


def parse_phone(value: Any, field: PhoneField) -> str | None:
    text = str(value or "").strip()
    if not text:
        return None
    digits = re.sub(r"\D", "", text)
    if not digits:
        return None
    if text.lstrip().startswith(("+", "00")):
        digits = digits[2:] if text.lstrip().startswith("00") else digits
        return f"+{digits}"
    from app.importers.jotform.convert import CALLING_CODES

    code = next((c for c, iso in CALLING_CODES.items() if iso == field.default_country), None)
    if code and not digits.startswith(code):
        return f"+{code}{digits}"
    return f"+{digits}"


def _links(value: Any) -> list[str]:
    return [u for u in re.split(r"[\s,]+", str(value or "")) if re.match(r"^https?://", u)]


def _choice(field: SingleChoiceField | MultiChoiceField, text: str) -> str | None:
    for o in field.options:
        if o.label == text or o.value == text:
            return o.value
    low = text.lower()
    for o in field.options:
        if o.label.lower() == low or o.value.lower() == low:
            return o.value
    return None


@dataclass
class RowResult:
    data: dict[str, Any] = field(default_factory=dict)
    links: dict[str, list[str]] = field(default_factory=dict)
    submitted_at: datetime | None = None
    external_id: str | None = None
    ip: str | None = None
    # (column index, problem code, cell text)
    problems: list[tuple[int, str, str]] = field(default_factory=list)


def convert_value(field: FieldElement, value: Any, order: str) -> tuple[Any, str | None]:
    """(answer, problem code). ``None`` answer with no problem means "blank"."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None, None
    text = value.strip() if isinstance(value, str) else str(value)
    out: Any
    if isinstance(field, ShortTextField | LongTextField | HiddenField):
        if isinstance(field, ShortTextField):
            text = " ".join(text.split())
        out = text
    elif isinstance(field, EmailField):
        out = text
    elif isinstance(field, PhoneField):
        out = parse_phone(text, field)
    elif isinstance(field, NumberField):
        out = parse_number(value)
    elif isinstance(field, DateField):
        day = parse_date(value, order)
        out = day.isoformat() if day else None
    elif isinstance(field, TimeField):
        out = parse_time(value)
    elif isinstance(field, SingleChoiceField):
        out = _choice(field, text)
        if out is None and field.allow_other and field.display != "dropdown":
            out = text
    elif isinstance(field, MultiChoiceField):
        parts = [p.strip() for p in re.split(r"\r?\n", text) if p.strip()]
        if len(parts) == 1 and _choice(field, parts[0]) is None:
            parts = [p.strip() for p in text.split(", ") if p.strip()]
        picked: list[str] = []
        others: list[str] = []
        for p in parts:
            v = _choice(field, p)
            if v is not None:
                if v not in picked:
                    picked.append(v)
            else:
                others.append(p)
        if others and field.allow_other:
            picked.append(", ".join(others))
            others = []
        out = None if others else picked
    elif isinstance(field, YesNoField):
        low = text.lower()
        out = True if low in YES else False if low in NO else None
    elif isinstance(field, RatingField):
        number = parse_number(value)
        out = (
            int(number) if isinstance(number, int | float) and float(number).is_integer() else None
        )
    else:  # pragma: no cover - file and signature are handled by the caller
        return None, "unsupported"
    if out is None:
        return None, "unreadable"
    problem = check_answer(field, out, False, date.today())
    if problem is not None and problem.code not in _ACCEPTED:
        return None, problem.code
    return out, None


def convert_row(
    row: list[Any],
    mapping: list[str],
    fields: dict[str, FieldElement],
    tz: tzinfo,
    order: str = "auto",
) -> RowResult:
    res = RowResult()
    for i, target in enumerate(mapping):
        if i >= len(row) or target == IGNORE:
            continue
        value = row[i]
        blank = value is None or (isinstance(value, str) and not value.strip())
        if target == "meta:submitted_at":
            if not blank:
                res.submitted_at = parse_datetime(value, tz, order)
                if res.submitted_at is None:
                    res.problems.append((i, "date", str(value)))
            continue
        if target == "meta:external_id":
            if not blank:
                text = str(value).strip()
                res.external_id = text[:-2] if re.fullmatch(r"\d+\.0", text) else text[:200]
            continue
        if target == "meta:ip":
            if not blank and re.fullmatch(r"[0-9A-Fa-f:.]{2,45}", str(value).strip()):
                res.ip = str(value).strip()
            continue
        if not target.startswith("field:") or blank:
            continue
        field_ = fields.get(target[6:])
        if field_ is None:
            continue
        if isinstance(field_, FileField | SignatureField):
            links = _links(value)
            if links:
                res.links.setdefault(field_.id, []).extend(links)
            else:
                res.problems.append((i, "no_link", str(value)[:200]))
            continue
        answer, problem = convert_value(field_, value, order)
        if problem:
            res.problems.append((i, problem, str(value)[:200]))
        elif answer is not None:
            res.data[field_.id] = answer
    return res


def jsonable(value: Any) -> Any:
    """A cell as JSON (dates as ISO text) for the stored sample rows."""
    if isinstance(value, datetime | date | time):
        return value.isoformat()
    if isinstance(value, int | float | str) or value is None:
        return value
    return str(value)
