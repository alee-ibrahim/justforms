"""Exports and uploaded images: respondent file names never become spreadsheet formulas, and
deeply nested SVGs are refused cleanly."""

from __future__ import annotations

import io
import uuid
from types import SimpleNamespace
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook

from app.branding.images import BrandImageError, clean_image
from app.config import Settings
from app.schemas.form import FileField
from app.submissions.export import Column, Context, cell_value
from app.uploads.service import clean_name


def _xlsx_cell(value: Any) -> Any:
    wb = Workbook(write_only=True)
    ws = wb.create_sheet()
    ws.append([value])
    buf = io.BytesIO()
    wb.save(buf)
    return load_workbook(io.BytesIO(buf.getvalue())).worksheets[0]["A1"]


@pytest.mark.parametrize(
    "name", ["=1+1", '=HYPERLINK("x","Open file")&"x.pdf', "+1.pdf", "@SUM(A1).pdf"]
)
def test_file_name_cells_are_never_formulas(name: str) -> None:
    el = FileField.model_validate({"kind": "field", "id": "doc", "type": "file", "label": "Doc"})
    ctx = Context(Settings(environment="test"), "en", {})
    value = [{"fileId": str(uuid.uuid4()), "name": clean_name(name)}]
    text = cell_value(Column("doc", "Doc", el), value, SimpleNamespace(id=uuid.uuid4()), ctx)  # type: ignore[arg-type]
    assert isinstance(text, str) and not text.startswith(("=", "+", "-", "@"))
    cell = _xlsx_cell(text)
    assert cell.data_type == "s"


def test_deeply_nested_svg_is_refused() -> None:
    head = '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 10 10">'
    deep = head + "<g>" * 5000 + "</g>" * 5000 + "</svg>"
    with pytest.raises(BrandImageError) as e:
        clean_image("logo", deep.encode())
    assert e.value.code == "svg_unsafe"
    ok = head + "<g>" * 10 + "<rect/>" + "</g>" * 10 + "</svg>"
    assert clean_image("logo", ok.encode()).mime == "image/svg+xml"
