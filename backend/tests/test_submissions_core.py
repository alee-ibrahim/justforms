"""Phase 3 submissions core: filters, masking, comments, bulk actions, deletes, retention,
the Excel export and the role x endpoint matrix (DESIGN 7, 15)."""

from __future__ import annotations

import io
import json
import time
import uuid
from datetime import UTC, datetime, timedelta
from typing import Any

import pytest
from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy import func, insert, select, text, update
from sqlalchemy.ext.asyncio import AsyncSession

from app.config import Settings
from app.integrity import verify_submission
from app.models import (
    AuditLog,
    Export,
    ExportStatus,
    FormPermission,
    FormRole,
    Submission,
    User,
)
from app.storage import Storage
from app.submissions import retention
from app.submissions.export import build_export, safe_text
from tests.conftest import RecordingQueue, login, make_client
from tests.helpers import Published, admin_client, field, publish

OPTS = [{"value": "civic", "label": "Civic"}, {"value": "tech", "label": "Technology"}]
INTERESTS = [{"value": "law", "label": "Law"}, {"value": "it", "label": "IT"}]

ELEMENTS = [
    field("name", label="Full name"),
    field("national_id", label="National ID", sensitive=True),
    field("age", "number", label="Age"),
    field("track", "single_choice", label="Track", options=OPTS),
    field("interests", "multi_choice", label="Interests", options=INTERESTS),
    field("start", "date", label="Start"),
    field("notes", "long_text", label="Notes"),
    {
        "kind": "repeater",
        "id": "dependents",
        "label": "Dependents",
        "elements": [field("dep_name", label="Name"), field("dep_id", label="ID", sensitive=True)],
    },
]

PEOPLE = [
    {
        "name": "Aminath Ali",
        "national_id": "A111567",
        "age": 19,
        "track": "civic",
        "interests": ["law"],
        "start": "2026-11-01",
        "notes": '=HYPERLINK("http://evil")',
        "dependents": [{"dep_name": "Aisha", "dep_id": "A900001"}],
    },
    {
        "name": "Ibrahim Naseer",
        "national_id": "A222903",
        "age": 31,
        "track": "tech",
        "interests": ["law", "it"],
        "start": "2027-01-15",
    },
    {"name": "Mariyam Hassan", "national_id": "A333221", "age": 24, "track": "tech"},
]


async def _seed(client: AsyncClient, me: dict[str, Any], language: str = "en") -> Published:
    pub = await publish(client, me, ELEMENTS, title="Intake", language=language)
    for person in PEOPLE:
        res = await client.post(f"{pub.base}/submissions", json=pub.submit_body(person))
        assert res.status_code == 201, res.text
    return pub


async def _list(client: AsyncClient, pub: Published, **params: Any) -> dict[str, Any]:
    if "filter" in params:
        params["filter"] = json.dumps(params["filter"])
    res = await client.get(f"/api/v1/forms/{pub.id}/submissions", params=params)
    assert res.status_code == 200, res.text
    return dict(res.json())


def names(page: dict[str, Any]) -> list[str]:
    return sorted(i["data"]["name"] for i in page["items"])


async def test_filters(
    client: AsyncClient, db: AsyncSession, settings: Settings, queue: RecordingQueue
) -> None:
    me = await admin_client(client, db, settings, queue)
    pub = await _seed(client, me)

    page = await _list(client, pub)
    assert page["total"] == 3
    assert names(await _list(client, pub, q="naseer")) == ["Ibrahim Naseer"]
    assert names(await _list(client, pub, q="technology")) == ["Ibrahim Naseer", "Mariyam Hassan"]
    # Sensitive answers are not searchable.
    assert (await _list(client, pub, q="A222903"))["total"] == 0

    def f(**kw: Any) -> list[dict[str, Any]]:
        return [kw]

    assert names(await _list(client, pub, filter=f(field="age", op="gte", value=24))) == [
        "Ibrahim Naseer",
        "Mariyam Hassan",
    ]
    assert names(await _list(client, pub, filter=f(field="age", op="lte", value=19))) == [
        "Aminath Ali"
    ]
    assert names(await _list(client, pub, filter=f(field="track", op="eq", value="civic"))) == [
        "Aminath Ali"
    ]
    assert (
        len((await _list(client, pub, filter=f(field="track", op="neq", value="civic")))["items"])
        == 2
    )
    assert names(
        await _list(client, pub, filter=f(field="interests", op="includes", value="it"))
    ) == ["Ibrahim Naseer"]
    assert names(await _list(client, pub, filter=f(field="name", op="contains", value="HASS"))) == [
        "Mariyam Hassan"
    ]
    assert names(
        await _list(client, pub, filter=f(field="start", op="gte", value="2027-01-01"))
    ) == ["Ibrahim Naseer"]
    assert names(await _list(client, pub, filter=f(field="interests", op="unanswered"))) == [
        "Mariyam Hassan"
    ]
    assert len((await _list(client, pub, filter=f(field="interests", op="answered")))["items"]) == 2
    both = [
        {"field": "track", "op": "eq", "value": "tech"},
        {"field": "age", "op": "gte", "value": 30},
    ]
    assert names(await _list(client, pub, filter=both)) == ["Ibrahim Naseer"]

    today = datetime.now(UTC).date().isoformat()
    assert (await _list(client, pub, **{"from": today}))["total"] == 3
    assert (await _list(client, pub, to="2020-01-01"))["total"] == 0
    assert (await _list(client, pub, version=1))["total"] == 3
    assert (await _list(client, pub, version=2))["total"] == 0

    for bad in ("not json", json.dumps([{"field": "Age!", "op": "eq", "value": 1}])):
        res = await client.get(f"/api/v1/forms/{pub.id}/submissions", params={"filter": bad})
        assert res.status_code == 422
    res = await client.get(
        f"/api/v1/forms/{pub.id}/submissions",
        params={"filter": json.dumps([{"field": "age", "op": "gte"}])},
    )
    assert res.status_code == 422


async def test_masking_reveal_comments(
    app: Any, client: AsyncClient, db: AsyncSession, settings: Settings, queue: RecordingQueue
) -> None:
    me = await admin_client(client, db, settings, queue)
    pub = await _seed(client, me)
    page = await _list(client, pub)
    row = next(i for i in page["items"] if i["data"]["name"] == "Aminath Ali")
    assert row["data"]["national_id"] == "••••567"
    assert row["data"]["dependents"][0]["dep_id"] == "••••001"

    sid = row["id"]
    detail = (await client.get(f"/api/v1/submissions/{sid}")).json()
    assert detail["data"]["national_id"] == "••••567" and detail["canDelete"] is True
    res = await client.post(f"/api/v1/submissions/{sid}/reveal", json={"fieldId": "national_id"})
    assert res.json() == {"value": "A111567"}
    res = await client.post(
        f"/api/v1/submissions/{sid}/reveal", json={"fieldId": "dep_id", "row": 0}
    )
    assert res.json() == {"value": "A900001"}
    res = await client.post(f"/api/v1/submissions/{sid}/reveal", json={"fieldId": "name"})
    assert res.status_code == 404

    res = await client.post(f"/api/v1/submissions/{sid}/comments", json={"text": "  Looks fine.  "})
    assert res.status_code == 201
    events = res.json()["events"]
    assert [e["type"] for e in events[:3]] == ["comment", "viewed", "viewed"]
    assert events[0]["payload"] == {"text": "Looks fine."} and events[0]["actorName"]
    res = await client.post(f"/api/v1/submissions/{sid}/comments", json={"text": "   "})
    assert res.status_code == 422
    sub = await db.get(Submission, uuid.UUID(sid))
    assert sub is not None and (await verify_submission(db, sub)).ok


async def test_detail_by_reference(
    client: AsyncClient, db: AsyncSession, settings: Settings, queue: RecordingQueue
) -> None:
    me = await admin_client(client, db, settings, queue)
    pub = await _seed(client, me)
    page = await _list(client, pub)
    row = next(i for i in page["items"] if i["data"]["name"] == "Ibrahim Naseer")
    res = await client.get(f"/api/v1/forms/{pub.id}/submissions/seq/{row['seq']}")
    assert res.status_code == 200, res.text
    detail = res.json()
    assert detail["id"] == row["id"] and detail["seq"] == row["seq"]
    assert detail["data"]["national_id"] == "••••903" and detail["schema"]["pages"]
    assert (await client.get(f"/api/v1/forms/{pub.id}/submissions/seq/99")).status_code == 404
    assert (await client.get(f"/api/v1/forms/{pub.id}/submissions/seq/0")).status_code == 422
    # Deleted submissions are gone by reference too; other forms' references don't leak.
    assert (await client.delete(f"/api/v1/submissions/{row['id']}")).status_code == 204
    res = await client.get(f"/api/v1/forms/{pub.id}/submissions/seq/{row['seq']}")
    assert res.status_code == 404
    other = await publish(client, me, ELEMENTS[:1], title="Other")
    assert (await client.get(f"/api/v1/forms/{other.id}/submissions/seq/1")).status_code == 404


async def test_card_fields(
    app: Any, client: AsyncClient, db: AsyncSession, settings: Settings, queue: RecordingQueue
) -> None:
    me = await admin_client(client, db, settings, queue)
    pub = await publish(client, me, ELEMENTS[:4], title="Cards")
    url = f"/api/v1/forms/{pub.id}"
    assert (await client.get(url)).json()["cardFields"] is None
    res = await client.patch(url, json={"cardFields": ["track", "name", "track"]})
    assert res.status_code == 200 and res.json()["cardFields"] == ["track", "name"]
    # Other settings stay; an empty list goes back to the automatic choice.
    res = await client.patch(url, json={"retentionDays": 30})
    assert res.json()["cardFields"] == ["track", "name"]
    res = await client.patch(url, json={"cardFields": []})
    assert res.json()["cardFields"] is None and res.json()["retentionDays"] == 30
    assert (await client.patch(url, json={"cardFields": ["Bad key"]})).status_code == 422
    too_many = [f"q{i}" for i in range(7)]
    assert (await client.patch(url, json={"cardFields": too_many})).status_code == 422

    # Editors choose the card fields; viewers can't.
    for role, expected in (("editor", 200), ("viewer", 403)):
        user = User(email=f"cards-{role}@example.org")
        db.add(user)
        await db.flush()
        db.add(FormPermission(form_id=pub.id, user_id=user.id, role=FormRole(role)))
        await db.commit()
        async with make_client(app) as other:
            await login(other, queue, user.email)
            res = await other.patch(url, json={"cardFields": ["age"]})
        assert res.status_code == expected, (role, res.text)


async def test_bulk_and_delete(
    client: AsyncClient,
    db: AsyncSession,
    settings: Settings,
    queue: RecordingQueue,
    file_storage: Storage,
) -> None:
    me = await admin_client(client, db, settings, queue)
    pub = await _seed(client, me)
    ids = [i["id"] for i in (await _list(client, pub))["items"]]
    url = f"/api/v1/forms/{pub.id}/submissions/bulk"

    res = await client.post(url, json={"ids": ids, "action": "status", "status": "rejected"})
    assert res.status_code == 422 and res.json()["detail"] == "reason_required"
    res = await client.post(
        url, json={"ids": ids[:2], "action": "status", "status": "rejected", "reason": "Late"}
    )
    assert res.json() == {"changed": 2}
    res = await client.post(
        url, json={"ids": ids, "action": "status", "status": "rejected", "reason": "x"}
    )
    assert res.json() == {"changed": 1}  # two were already rejected
    stats = (await client.get(f"/api/v1/forms/{pub.id}/submissions/stats")).json()
    assert stats["rejected"] == 3

    # Ids of another form are ignored.
    other = await publish(client, me, [field("q")], title="Other")
    await client.post(f"{other.base}/submissions", json=other.submit_body({"q": "x"}))
    other_id = (await _list(client, other))["items"][0]["id"]
    res = await client.post(url, json={"ids": [other_id], "action": "status", "status": "approved"})
    assert res.json() == {"changed": 0}

    # Soft delete hides at once; the nightly job erases after 30 days.
    res = await client.post(url, json={"ids": ids[:1], "action": "delete"})
    assert res.json() == {"changed": 1}
    assert (await _list(client, pub))["total"] == 2
    assert (await client.get(f"/api/v1/submissions/{ids[0]}")).status_code == 404
    assert (await client.get(f"/api/v1/forms/{pub.id}/submissions/stats")).json()["all"] == 2
    forms = (await client.get(f"/api/v1/orgs/{me['orgs'][0]['id']}/forms")).json()
    intake = next(f for f in forms if f["id"] == pub.id)
    assert intake["submissionCount"] == 2 and intake["newCount"] == 0

    assert await retention.purge(db, file_storage) == 0
    await db.execute(
        update(Submission)
        .where(Submission.id == uuid.UUID(ids[0]))
        .values(deleted_at=datetime.now(UTC) - timedelta(days=31))
    )
    await db.commit()
    assert await retention.purge(db, file_storage) == 1
    assert await db.get(Submission, uuid.UUID(ids[0])) is None
    tomb = await db.scalar(select(AuditLog).where(AuditLog.action == "submission.erased"))
    assert tomb is not None and tomb.meta["reason"] == "deleted"

    # Owner hard delete.
    res = await client.delete(f"/api/v1/submissions/{ids[1]}?hard=true")
    assert res.status_code == 204
    assert (
        await db.scalar(
            select(func.count()).select_from(Submission).where(Submission.id == uuid.UUID(ids[1]))
        )
        == 0
    )


async def test_retention_days(
    client: AsyncClient,
    db: AsyncSession,
    settings: Settings,
    queue: RecordingQueue,
    file_storage: Storage,
) -> None:
    me = await admin_client(client, db, settings, queue)
    pub = await _seed(client, me)
    res = await client.patch(f"/api/v1/forms/{pub.id}", json={"retentionDays": 7})
    assert res.status_code == 200 and res.json()["retentionDays"] == 7
    await db.execute(
        text("UPDATE submissions SET submitted_at = now() - interval '8 days' WHERE seq = 1")
    )
    await db.commit()
    assert await retention.purge(db, file_storage) == 1
    assert (await _list(client, pub))["total"] == 2
    res = await client.patch(f"/api/v1/forms/{pub.id}", json={"retentionDays": 0})
    assert res.json()["retentionDays"] is None


# -- export ----------------------------------------------------------------------------------


async def _export(
    client: AsyncClient,
    db: AsyncSession,
    settings: Settings,
    queue: RecordingQueue,
    storage: Storage,
    pub: Published,
    body: dict[str, Any] | None = None,
) -> Any:
    res = await client.post(f"/api/v1/forms/{pub.id}/exports", json=body or {})
    assert res.status_code == 201, res.text
    eid = res.json()["id"]
    assert ("export_submissions", {"export_id": eid}) in queue.jobs
    assert await build_export(db, storage, settings, uuid.UUID(eid))
    info = (await client.get(f"/api/v1/exports/{eid}")).json()
    assert info["status"] == "ready" and info["downloadUrl"], info
    res = await client.get(info["downloadUrl"])
    assert res.status_code == 200
    return load_workbook(io.BytesIO(res.content)), info


def rows(ws: Any) -> list[list[Any]]:
    return [list(r) for r in ws.iter_rows(values_only=True)]


async def test_export_xlsx(
    client: AsyncClient,
    db: AsyncSession,
    settings: Settings,
    queue: RecordingQueue,
    file_storage: Storage,
) -> None:
    me = await admin_client(client, db, settings, queue)
    pub = await _seed(client, me)
    wb, info = await _export(client, db, settings, queue, file_storage, pub)
    assert info["rowCount"] == 3 and info["rowsDone"] == 3
    assert wb.sheetnames == ["Submissions", "Dependents"]
    sheet = rows(wb["Submissions"])
    assert sheet[0][:6] == [
        "#",
        "Submitted at",
        "Status",
        "Version",
        "Language",
        "Respondent email",
    ]
    assert sheet[0][6:] == [
        "Full name (name)",
        "National ID (national_id)",
        "Age (age)",
        "Track (track)",
        "Interests (interests)",
        "Start (start)",
        "Notes (notes)",
    ]
    first = sheet[1]
    assert first[0] == 1 and first[2] == "New" and first[3] == 1 and first[4] == "English"
    assert first[5] is None  # public form: no confirmed respondent address
    assert first[6:13] == [
        "Aminath Ali",
        "A111567",  # editors and owners export sensitive answers
        19,
        "Civic",
        "Law",
        "2026-11-01",
        '\'=HYPERLINK("http://evil")',  # formula injection neutralised
    ]
    assert sheet[2][10] == "Law; IT"
    assert rows(wb["Dependents"]) == [
        ["Submission #", "Row", "Name (dep_name)", "ID (dep_id)"],
        [1, 1, "Aisha", "A900001"],
    ]
    assert wb["Submissions"].sheet_view.rightToLeft in (None, False)

    # Only the selection; filters.
    ids = [i["id"] for i in (await _list(client, pub))["items"]]
    wb, info = await _export(client, db, settings, queue, file_storage, pub, {"ids": ids[:1]})
    assert info["rowCount"] == 1
    wb, info = await _export(
        client,
        db,
        settings,
        queue,
        file_storage,
        pub,
        {"fields": [{"field": "track", "op": "eq", "value": "tech"}]},
    )
    assert info["rowCount"] == 2

    # A new version: removed questions move to the end, new ones follow the latest order.
    draft = (await client.get(f"/api/v1/forms/{pub.id}/draft")).json()
    res = await client.patch(
        f"/api/v1/forms/{pub.id}/draft",
        json={
            "baseRevision": draft["revision"],
            "patch": [
                {"op": "remove", "path": "/pages/0/elements/6"},
                {
                    "op": "add",
                    "path": "/pages/0/elements/1",
                    "value": field("email", "email", label="Email"),
                },
            ],
        },
    )
    assert res.status_code == 200, res.text
    assert (await client.post(f"/api/v1/forms/{pub.id}/publish")).status_code == 200
    wb, _info = await _export(client, db, settings, queue, file_storage, pub)
    header = rows(wb["Submissions"])[0]
    assert header[6:8] == ["Full name (name)", "Email (email)"]
    assert header[-1] == "Notes (version 1) (notes)"


async def test_export_dv_and_viewer_masking(
    app: Any,
    client: AsyncClient,
    db: AsyncSession,
    settings: Settings,
    queue: RecordingQueue,
    file_storage: Storage,
) -> None:
    me = await admin_client(client, db, settings, queue)
    pub = await _seed(client, me, language="dv")
    viewer = User(email="viewer@example.org")
    db.add(viewer)
    await db.flush()
    db.add(FormPermission(form_id=pub.id, user_id=viewer.id, role=FormRole.VIEWER))
    await db.commit()
    async with make_client(app) as other:
        await login(other, queue, "viewer@example.org")
        wb, _ = await _export(other, db, settings, queue, file_storage, pub)
        # Someone else cannot fetch this export.
        eid = (await db.scalars(select(Export).order_by(Export.created_at.desc()))).first()
        assert eid is not None
    assert (await client.get(f"/api/v1/exports/{eid.id}")).status_code == 404
    assert wb["Submissions"].sheet_view.rightToLeft is True
    sheet = rows(wb["Submissions"])
    assert sheet[1][4] == "Dhivehi" and sheet[1][7] == "••••567"
    assert rows(wb["Dependents"])[1][3] == "••••001"


def test_safe_text() -> None:
    for s in ("=1+1", "+1", "-1", "@SUM(A1)", "\tx"):
        assert safe_text(s) == "'" + s
    assert safe_text("Aminath") == "Aminath"


async def test_export_ten_thousand_rows(
    client: AsyncClient,
    db: AsyncSession,
    settings: Settings,
    queue: RecordingQueue,
    file_storage: Storage,
) -> None:
    me = await admin_client(client, db, settings, queue)
    pub = await publish(client, me, ELEMENTS, title="Big", language="dv")
    version_id = uuid.UUID(pub.live["versionId"])
    now = datetime.now(UTC)
    await db.execute(
        insert(Submission),
        [
            {
                "id": uuid.uuid4(),
                "form_id": uuid.UUID(pub.id),
                "form_version_id": version_id,
                "seq": i + 1,
                "language": "dv",
                "data": {
                    "name": f"ޢާއިޝަތު {i}",
                    "age": 18 + i % 20,
                    "track": "tech",
                    "dependents": [{"dep_name": "ޢަލީ"}, {"dep_name": "ހަސަން"}],
                },
                "submitted_at": now,
                "search_text": "",
            }
            for i in range(10_000)
        ],
    )
    await db.commit()
    started = time.monotonic()
    wb, info = await _export(client, db, settings, queue, file_storage, pub)
    elapsed = time.monotonic() - started
    assert info["rowCount"] == 10_000
    assert wb["Submissions"].max_row == 10_001
    assert wb["Dependents"].max_row == 20_001
    assert rows(wb["Submissions"])[10_000][6] == "ޢާއިޝަތު 9999"
    assert elapsed < 60, elapsed


# -- role x endpoint matrix ------------------------------------------------------------------

ROLES = ["viewer", "editor", "owner", "outsider"]

# endpoint -> (method, path, body, minimum role)
MATRIX: dict[str, tuple[str, str, Any, str]] = {
    "list": ("GET", "/api/v1/forms/{form}/submissions", None, "viewer"),
    "stats": ("GET", "/api/v1/forms/{form}/submissions/stats", None, "viewer"),
    "detail": ("GET", "/api/v1/submissions/{sub}", None, "viewer"),
    "by_reference": ("GET", "/api/v1/forms/{form}/submissions/seq/1", None, "viewer"),
    "verify": ("GET", "/api/v1/submissions/{sub}/verify", None, "viewer"),
    "comment": ("POST", "/api/v1/submissions/{sub}/comments", {"text": "hi"}, "viewer"),
    "request_pdf": ("POST", "/api/v1/submissions/{sub}/pdf", None, "viewer"),
    "export": ("POST", "/api/v1/forms/{form}/exports", {}, "viewer"),
    "status": ("PATCH", "/api/v1/submissions/{sub}", {"status": "in_review"}, "editor"),
    "reveal": ("POST", "/api/v1/submissions/{sub}/reveal", {"fieldId": "national_id"}, "editor"),
    "bulk_status": (
        "POST",
        "/api/v1/forms/{form}/submissions/bulk",
        {"ids": ["{sub}"], "action": "status", "status": "approved"},
        "editor",
    ),
    "bulk_delete": (
        "POST",
        "/api/v1/forms/{form}/submissions/bulk",
        {"ids": ["{sub}"], "action": "delete"},
        "owner",
    ),
    "delete": ("DELETE", "/api/v1/submissions/{sub}", None, "owner"),
}
RANK = {"viewer": 1, "editor": 2, "owner": 3, "outsider": 0}


@pytest.mark.parametrize("role", ROLES)
@pytest.mark.parametrize("endpoint", list(MATRIX))
async def test_submission_role_matrix(
    role: str,
    endpoint: str,
    app: Any,
    client: AsyncClient,
    db: AsyncSession,
    settings: Settings,
    queue: RecordingQueue,
) -> None:
    me = await admin_client(client, db, settings, queue)
    pub = await publish(client, me, ELEMENTS[:2], title="Matrix")
    res = await client.post(
        f"{pub.base}/submissions", json=pub.submit_body({"name": "A", "national_id": "A1"})
    )
    assert res.status_code == 201
    sid = (await _list(client, pub))["items"][0]["id"]

    user = User(email=f"{role}@example.org")
    db.add(user)
    await db.flush()
    if role != "outsider":
        db.add(FormPermission(form_id=pub.id, user_id=user.id, role=FormRole(role)))
    await db.commit()

    method, path, body, minimum = MATRIX[endpoint]
    path = path.format(form=pub.id, sub=sid)
    if body is not None:
        body = json.loads(json.dumps(body).replace("{sub}", sid))
    async with make_client(app) as other:
        await login(other, queue, user.email)
        res = await other.request(method, path, json=body)
    allowed = RANK[role] >= RANK[minimum]
    if allowed:
        assert res.status_code in (200, 201, 204), (endpoint, role, res.status_code, res.text)
    else:
        assert res.status_code == (404 if role == "outsider" else 403), (
            endpoint,
            role,
            res.status_code,
        )
    sub = await db.get(Submission, uuid.UUID(sid))
    assert sub is not None
    await db.refresh(sub)
    if endpoint == "status":
        assert (sub.status.value == "in_review") is allowed
    if endpoint in ("delete", "bulk_delete"):
        assert (sub.deleted_at is not None) is allowed


async def test_export_progress_states(
    client: AsyncClient, db: AsyncSession, settings: Settings, queue: RecordingQueue
) -> None:
    me = await admin_client(client, db, settings, queue)
    pub = await _seed(client, me)
    res = await client.post(f"/api/v1/forms/{pub.id}/exports", json={})
    info = res.json()
    assert info["status"] == "pending" and info["downloadUrl"] is None
    assert (await client.get(f"/api/v1/exports/{info['id']}/download")).status_code == 404
    row = await db.get(Export, uuid.UUID(info["id"]))
    assert row is not None and row.status == ExportStatus.PENDING
