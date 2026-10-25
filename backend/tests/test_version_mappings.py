"""Version field mappings: an older version's fixed questions as rows of a newer repeater,
in the form's settings and the Excel export (DESIGN 15.2)."""

from __future__ import annotations

import io
import uuid
from typing import Any

from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from app.config import Settings
from app.schemas.form import Form
from app.storage import Storage
from app.submissions import version_map
from app.submissions.export import build_export
from tests.conftest import RecordingQueue
from tests.helpers import admin_client, field, publish

V1 = [
    field("org", label="Organisation", required=True),
    *(
        f
        for n in (1, 2, 3)
        for f in (
            field(f"member{n}_name", label="Full name"),
            field(f"member{n}_role", label="Role"),
        )
    ),
    field("date", "date", label="Date"),
]
MEMBERS = {
    "kind": "repeater",
    "id": "members",
    "label": "Members",
    "minRows": 1,
    "elements": [field("member_name", label="Full name"), field("member_role", label="Role")],
}
MAPPING = {
    "repeater": "members",
    "rows": [
        {"member_name": f"member{n}_name", "member_role": f"member{n}_role"} for n in (1, 2, 3)
    ],
}


def form(elements: list[dict[str, Any]]) -> Form:
    return Form.model_validate(
        {
            "schemaVersion": "1",
            "settings": {"title": "T"},
            "pages": [{"id": "p", "elements": elements}],
        }
    )


def test_mapped_rows_skip_empty_slots() -> None:
    [m] = version_map.parse([MAPPING])
    assert m.sources == {f"member{n}_{k}" for n in (1, 2, 3) for k in ("name", "role")}
    data = {"member1_name": "A", "member1_role": "Chair", "member2_name": "", "member3_role": "X"}
    assert version_map.mapped_rows(m, data) == [
        {"member_name": "A", "member_role": "Chair"},
        {"member_role": "X"},
    ]
    assert version_map.mapped_rows(m, {}) == []
    # Malformed settings are ignored, never fatal.
    assert version_map.parse("x") == [] and version_map.parse([{"rows": []}, 3]) == []


def test_problems() -> None:
    v1, v2 = form(V1), form([field("org"), MEMBERS])
    [m] = version_map.parse([MAPPING])
    assert version_map.problems([m], v2, [v1]) == []
    bad = version_map.parse(
        [
            {"repeater": "org", "rows": [{"x": "member1_name"}]},
            {"repeater": "members", "rows": [{"nope": "member1_name", "member_role": "ghost"}]},
            {"repeater": "members", "rows": [{"member_name": "org"}]},
        ]
    )
    found = version_map.problems(bad, v2, [v1])
    assert "org: not a repeater of this form" in found
    assert "members: nope is not a question in the repeater" in found
    assert "members: ghost is not a question of a published version" in found
    assert "members: org is still a question of the form" in found
    assert "members: mapped twice" in found


async def _export(
    client: AsyncClient, db: AsyncSession, settings: Settings, storage: Storage, fid: str
) -> Any:
    res = await client.post(f"/api/v1/forms/{fid}/exports", json={})
    assert res.status_code == 201, res.text
    assert await build_export(db, storage, settings, uuid.UUID(res.json()["id"]))
    info = (await client.get(f"/api/v1/exports/{res.json()['id']}")).json()
    data = (await client.get(info["downloadUrl"])).content
    wb = load_workbook(io.BytesIO(data))
    return {ws.title: [list(r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets}


async def test_export_reads_old_slots_as_repeater_rows(
    client: AsyncClient,
    db: AsyncSession,
    settings: Settings,
    queue: RecordingQueue,
    file_storage: Storage,
) -> None:
    me = await admin_client(client, db, settings, queue)
    pub = await publish(client, me, V1, title="Network")
    old = [
        {
            "org": "One",
            "member1_name": "Aminath",
            "member1_role": "Chair",
            "member2_name": "Ali",
            "member2_role": "Member",
            "date": "2026-10-01",
        },
        {"org": "Two", "member1_name": "Hassan", "member1_role": "Chair"},
    ]
    for data in old:
        res = await client.post(f"{pub.base}/submissions", json=pub.submit_body(data))
        assert res.status_code == 201, res.text

    url = f"/api/v1/forms/{pub.id}"
    # Before version 2 there is no repeater to map onto.
    res = await client.patch(url, json={"versionMappings": [MAPPING]})
    assert res.status_code == 422 and res.json()["detail"]["code"] == "bad_version_mapping"

    draft = (await client.get(f"{url}/draft")).json()
    res = await client.patch(
        f"{url}/draft",
        json={
            "baseRevision": draft["revision"],
            "patch": [{"op": "replace", "path": "/pages/0/elements", "value": [V1[0], MEMBERS]}],
        },
    )
    assert res.status_code == 200, res.text
    assert (await client.post(f"{url}/publish")).status_code == 200
    live = (await client.get(pub.base)).json()
    res = await client.post(
        f"{pub.base}/submissions",
        json={
            "versionId": live["versionId"],
            "uploadNonce": live["uploadNonce"],
            "data": {
                "org": "Three",
                "members": [{"member_name": "Mariyam", "member_role": "Chair"}],
            },
        },
    )
    assert res.status_code == 201, res.text

    # Without a mapping: the old slots are columns, labelled with their version.
    sheets = await _export(client, db, settings, file_storage, pub.id)
    assert "Full name (version 1) (member1_name)" in sheets["Submissions"][0]
    assert len(sheets["Members"]) == 2  # header + version 2's row

    res = await client.patch(url, json={"versionMappings": [MAPPING]})
    assert res.status_code == 200, res.text
    assert res.json()["versionMappings"] == [MAPPING]
    sheets = await _export(client, db, settings, file_storage, pub.id)
    head = sheets["Submissions"][0]
    assert head[6:] == ["Organisation (org)", "Date (version 1) (date)"]
    assert [r[6:] for r in sheets["Submissions"][1:]] == [
        ["One", "2026-10-01"],
        ["Two", None],
        ["Three", None],
    ]
    assert sheets["Members"] == [
        ["Submission #", "Row", "Full name (member_name)", "Role (member_role)"],
        [1, 1, "Aminath", "Chair"],
        [1, 2, "Ali", "Member"],
        [2, 1, "Hassan", "Chair"],
        [3, 1, "Mariyam", "Chair"],
    ]

    # An empty list removes the mapping.
    res = await client.patch(url, json={"versionMappings": []})
    assert res.json()["versionMappings"] == []
    assert (await client.get(url)).json()["versionMappings"] == []
