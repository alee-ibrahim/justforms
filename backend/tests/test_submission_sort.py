"""Sorting the submissions table by any column (DESIGN 15.1): server-side, every page in
order, empty answers last, unknown columns refused, exports in the table's order."""

from __future__ import annotations

import io
import uuid
from typing import Any

from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from app.config import Settings
from app.storage import Storage
from app.submissions.export import build_export
from tests.conftest import RecordingQueue
from tests.helpers import Published, admin_client, field, publish

OPTS = [
    {"value": "low", "label": "Zebra"},
    {"value": "mid", "label": "Apple"},
    {"value": "high", "label": "Mango"},
]
TAGS = [{"value": "a", "label": "A"}, {"value": "b", "label": "B"}, {"value": "c", "label": "C"}]

ELEMENTS = [
    field("name", label="Name"),
    field("age", "number", label="Age"),
    field("level", "single_choice", label="Level", options=OPTS, allowOther=True),
    field("tags", "multi_choice", label="Tags", options=TAGS),
    field("start", "date", label="Start"),
    field("phone", "phone", label="Phone"),
    field("ok", "yes_no", label="OK"),
    field("doc", "file", label="Document"),
]

ROWS: list[dict[str, Any]] = [
    {
        "name": "beta",
        "age": 9,
        "level": "high",
        "tags": ["c"],
        "start": "2026-03-01",
        "phone": "+9607771234",
        "ok": True,
    },
    {
        "name": "Alpha",
        "age": 100,
        "level": "low",
        "tags": ["b", "a"],
        "start": "2025-12-31",
        "phone": "+4420123456",
        "ok": False,
    },
    {"name": "ހަސަން", "age": 25.5, "level": "Something else", "start": "2026-01-15"},
    {"name": "ޢަލީ", "level": "mid", "tags": ["b"], "phone": "+9609990000"},
    {},
]


async def _seed(client: AsyncClient, me: dict[str, Any]) -> Published:
    pub = await publish(client, me, ELEMENTS, title="Sorting")
    for row in ROWS:
        res = await client.post(f"{pub.base}/submissions", json=pub.submit_body(row))
        assert res.status_code == 201, res.text
    return pub


async def _seqs(client: AsyncClient, pub: Published, sort: str, dir_: str, **kw: Any) -> list[int]:
    out: list[int] = []
    cursor = None
    while True:
        params: dict[str, Any] = {"sort": sort, "dir": dir_, "limit": 2, **kw}
        if cursor is not None:
            params["cursor"] = cursor
        res = await client.get(f"/api/v1/forms/{pub.id}/submissions", params=params)
        assert res.status_code == 200, res.text
        page = res.json()
        out += [i["seq"] for i in page["items"]]
        cursor = page["nextCursor"]
        if cursor is None:
            return out


async def test_sort_columns(
    client: AsyncClient, db: AsyncSession, settings: Settings, queue: RecordingQueue
) -> None:
    me = await admin_client(client, db, settings, queue)
    pub = await _seed(client, me)

    assert await _seqs(client, pub, "_seq", "asc") == [1, 2, 3, 4, 5]
    assert await _seqs(client, pub, "_seq", "desc") == [5, 4, 3, 2, 1]
    assert await _seqs(client, pub, "_submitted", "asc") == [1, 2, 3, 4, 5]
    # Text: case-insensitive alphabetical, Latin before Thaana, Thaana in alphabet order
    # (haa before ainu); empty last both ways.
    assert await _seqs(client, pub, "name", "asc") == [2, 1, 3, 4, 5]
    assert await _seqs(client, pub, "name", "desc") == [4, 3, 1, 2, 5]
    # Numbers numerically (100 after 25.5), empty last.
    # Ties (here: the empty ones) go newest first.
    assert await _seqs(client, pub, "age", "asc") == [1, 3, 2, 5, 4]
    assert await _seqs(client, pub, "age", "desc") == [2, 3, 1, 5, 4]
    # Choices by option order (low, mid, high), "Other" answers after the options.
    assert await _seqs(client, pub, "level", "asc") == [2, 4, 1, 3, 5]
    # Multiple choice by the first option chosen in option order.
    assert await _seqs(client, pub, "tags", "asc") == [2, 4, 1, 5, 3]
    # Dates chronologically.
    assert await _seqs(client, pub, "start", "asc") == [2, 3, 1, 5, 4]
    # Phones by digits.
    assert await _seqs(client, pub, "phone", "asc") == [2, 1, 4, 5, 3]
    # Yes/no: No, then Yes.
    assert await _seqs(client, pub, "ok", "asc") == [2, 1, 5, 4, 3]
    # Status order, filters still apply.
    assert await _seqs(client, pub, "_status", "asc", q="alpha") == [2]

    # Columns that are not answers of this form, files and garbage are refused.
    for bad in ("doc", "nope", "_secret", "name;drop", "DATA"):
        res = await client.get(f"/api/v1/forms/{pub.id}/submissions", params={"sort": bad})
        assert res.status_code == 422, bad


async def test_export_in_current_order(
    client: AsyncClient,
    db: AsyncSession,
    settings: Settings,
    queue: RecordingQueue,
    file_storage: Storage,
) -> None:
    me = await admin_client(client, db, settings, queue)
    pub = await _seed(client, me)
    for body, expected in (
        ({}, [1, 2, 3, 4, 5]),
        ({"sort": "age", "dir": "desc"}, [2, 3, 1, 5, 4]),
    ):
        res = await client.post(f"/api/v1/forms/{pub.id}/exports", json=body)
        assert res.status_code == 201, res.text
        export_id = res.json()["id"]
        assert await build_export(db, file_storage, settings, uuid.UUID(export_id))
        res = await client.get(f"/api/v1/exports/{export_id}/download")
        wb = load_workbook(io.BytesIO(res.content), read_only=True)
        ws = wb["Submissions"]
        seqs = [r[0] for r in ws.iter_rows(min_row=2, values_only=True)]
        assert seqs == expected
